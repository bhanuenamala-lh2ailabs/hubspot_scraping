# -*- coding: utf-8 -*-
"""IBBI liquidation / CIRP list (XLSX) + quarterly newsletter (PDF, aggregates only).

CURRENCY GUARD (spec §2c, binding): the list is point-in-time. For any candidate whose
shutdown_date is after the file's as-of date, absence from this file proves nothing, so the
IBBI bonus must not be applied. The as-of date is parsed from the filename/title, never
hardcoded, so a newer file drops in cleanly.
"""
from __future__ import annotations
import os, re, datetime
from typing import Optional
import openpyxl

from ..models import CIN_RE, norm_name, decode_cin
from ..db import upsert_entity, upsert_candidate, add_evidence

_MONTHS = ("january february march april may june july august september october "
           "november december").split()
# "as on 31st March, 2025" — the comma is optional and IBBI uses it, which is how a
# four-month error crept in on the first run.
_ASOF_TEXT = re.compile(
    r"as\s+on\s+(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]+)\s*,?\s*(\d{4})", re.I)
_ASOF_ISO = re.compile(r"(\d{4})-(\d{2})-(\d{2})")


def parse_as_of(text: str, allow_bare_date: bool = True) -> Optional[str]:
    """Pull the as-of date out of a title or filename. Never hardcoded.

    `allow_bare_date=False` for FILENAMES. A filename like
    "2025-07-30 16_45_39-2f9407c2....xlsx" carries the DOWNLOAD timestamp, not the
    reporting date — the real one ("as on 31st March, 2025") lives in the sheet title.
    Treating the download date as the as-of date silently widens the currency guard by
    four months, which is the exact failure §2c is written to prevent.
    """
    if not text: return None
    m = _ASOF_TEXT.search(text)
    if m:
        d, mon, y = m.groups()
        mon = mon.lower()
        for i, name in enumerate(_MONTHS, 1):
            if mon.startswith(name[:3]):
                return f"{y}-{i:02d}-{int(d):02d}"
    if allow_bare_date:
        m = _ASOF_ISO.search(text)
        if m: return "-".join(m.groups())
    return None


def _cells(ws, limit=None):
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if limit and i > limit: break
        yield [("" if c is None else str(c)).strip() for c in row]


def ingest(db, folder: str, log) -> dict:
    out = {"as_of": None, "rows": 0, "with_cin": 0, "voluntary": 0, "files": []}
    if not os.path.isdir(folder):
        log("  ibbi: folder missing — skipped"); return out
    xlsx = [f for f in os.listdir(folder) if f.lower().endswith((".xlsx", ".xls"))]
    pdfs = [f for f in os.listdir(folder) if f.lower().endswith(".pdf")]
    if not xlsx:
        log("  ibbi: no XLSX found — no company-level IBBI evidence this run")

    for fn in sorted(xlsx):
        path = os.path.join(folder, fn)
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            log(f"  ibbi PARSE FAIL {fn}: {e}"); continue
        ws = wb[wb.sheetnames[0]]

        as_of = parse_as_of(fn, allow_bare_date=False)
        header_row, headers = None, []
        preview = []
        for i, row in enumerate(_cells(ws, limit=25)):
            preview.append(" ".join(row))
            joined = " ".join(row).lower()
            if not as_of:
                a = parse_as_of(" ".join(row))     # sheet title: the authoritative one
                if a: as_of = a
            if header_row is None and ("name of the corporate" in joined
                                       or "corporate debtor" in joined
                                       or ("cin" in joined and "name" in joined)):
                header_row, headers = i, [c.lower() for c in row]
        if not as_of:
            log(f"  !! ibbi {fn}: NO as-of date found — IBBI bonus will be withheld entirely "
                f"for safety (cannot verify currency)")
        out["as_of"] = as_of or out["as_of"]
        out["files"].append({"file": fn, "as_of": as_of})

        # find the columns we need, by header text
        def col_of(*keys):
            for j, h in enumerate(headers):
                if any(k in h for k in keys): return j
            return None
        c_name = col_of("name of the corporate", "corporate debtor", "corporate person", "name")
        c_cin = col_of("cin", "llpin")
        n_rows = 0
        for row in _cells(ws):
            if header_row is not None and n_rows == 0 and " ".join(row).lower() == " ".join(headers):
                continue
            joined = " ".join(row)
            cin_m = CIN_RE.search(joined.upper())
            name = row[c_name].strip() if (c_name is not None and c_name < len(row)) else ""
            if not name and not cin_m: continue
            if name.lower() in ("", "name", "total", "nan"): continue
            if len(name) < 3 and not cin_m: continue
            n_rows += 1
            voluntary = "voluntary" in joined.lower()
            kind = "ibbi_voluntary" if voluntary else "ibbi"
            if voluntary: out["voluntary"] += 1
            cin = cin_m.group(0) if cin_m else None
            if cin:
                out["with_cin"] += 1
                d = decode_cin(cin)
                upsert_entity(db, cin, name or cin, norm_name(name or cin),
                              status="Under Liquidation (IBBI)", cls=d.get("company_class"),
                              dor=d.get("date_of_registration"),
                              state=d.get("registered_state"), source=f"ibbi:{fn}")
            if name:
                cid = upsert_candidate(db, name, norm_name(name), "ibbi")
                if cin:
                    db.execute("INSERT OR IGNORE INTO candidate_entity"
                               "(candidate_id,cin,match_score,match_method,is_confirmed) "
                               "VALUES(?,?,?,?,1)", (cid, cin, 100.0, "exact"))
                add_evidence(db, cid, kind, source_name="IBBI", as_of_date=as_of,
                             snippet=f"IBBI {'voluntary ' if voluntary else ''}liquidation list "
                                     f"(as on {as_of or 'unknown'})",
                             source_url="https://ibbi.gov.in/en/liquidation",
                             raw={"file": fn, "cin": cin})
        db.commit()
        out["rows"] += n_rows
        log(f"  ibbi {fn}: {n_rows} rows, as-of {as_of or 'UNKNOWN'}, {out['with_cin']} with CIN")

    for fn in sorted(pdfs):
        log(f"  ibbi newsletter {fn}: aggregates only — no entity extraction (per spec §2c)")
    return out


def bonus_allowed(shutdown_date: Optional[str], as_of: Optional[str]) -> bool:
    """The currency guard. No as-of date => never award the bonus."""
    if not as_of: return False
    if not shutdown_date: return True     # undated candidate: the list still confirms death
    try:
        sd = datetime.date.fromisoformat(shutdown_date[:10])
        ao = datetime.date.fromisoformat(as_of[:10])
    except Exception:
        return False
    return sd <= ao
