# -*- coding: utf-8 -*-
"""Shared .docx building helpers, tuned for clean Google Docs import.

Choices made for the Docs converter specifically:
  * real Heading 1/2/3 styles, so the Docs OUTLINE PANE populates (left-hand nav)
  * table style "Table Grid" — Docs keeps the borders; fancier Word styles get flattened
  * images inserted as PNG at an explicit width; Docs rejects SVG
  * shading applied per-cell via w:shd, which survives the conversion
  * no text boxes, no columns, no SmartArt — Docs drops or mangles all three
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK      = RGBColor(0x11, 0x1A, 0x26)
MUTED    = RGBColor(0x55, 0x60, 0x70)
BLUE     = RGBColor(0x1F, 0x4E, 0x87)
RED      = RGBColor(0xA3, 0x39, 0x2F)
AMBER    = RGBColor(0x8A, 0x4B, 0x00)
GREEN    = RGBColor(0x2F, 0x6B, 0x4F)

HDR_BG   = "E8EFF8"
NEW_BG   = "FDF0DD"
DEAD_BG  = "FBECEA"
LIVE_BG  = "E8F2EC"
BOX_BG   = "F4F6F8"


def new_doc(title, subtitle=None, intro=None):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"; st.font.size = Pt(10.5); st.font.color.rgb = INK
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for nm, sz, col in (("Heading 1", 17, BLUE), ("Heading 2", 13.5, INK), ("Heading 3", 11.5, INK)):
        s = d.styles[nm]
        s.font.name = "Calibri"; s.font.size = Pt(sz); s.font.bold = True; s.font.color.rgb = col
        s.paragraph_format.space_before = Pt(16 if nm == "Heading 1" else 11)
        s.paragraph_format.space_after = Pt(5)
    for sec in d.sections:
        sec.left_margin = sec.right_margin = Inches(0.85)
        sec.top_margin = sec.bottom_margin = Inches(0.75)
    t = d.add_heading(title, 0)
    t.runs[0].font.color.rgb = INK
    if subtitle:
        p = d.add_paragraph(); r = p.add_run(subtitle)
        r.font.size = Pt(12); r.font.color.rgb = MUTED; r.italic = True
    if intro:
        for para in intro:
            d.add_paragraph(para)
    return d


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _cell_text(cell, text, bold=False, size=9.5, color=None, mono=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
    for i, chunk in enumerate(str(text).split("\n")):
        para = p if i == 0 else cell.add_paragraph()
        para.paragraph_format.space_after = Pt(2)
        r = para.add_run(chunk)
        r.bold = bold; r.font.size = Pt(size)
        r.font.name = "Consolas" if mono else "Calibri"
        if color is not None: r.font.color.rgb = color


def table(doc, headers, rows, widths=None, mono_cols=(), row_bg=None):
    """row_bg: callable(row_index, row) -> hex or None"""
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        _cell_text(t.rows[0].cells[i], h, bold=True, size=9)
        _shade(t.rows[0].cells[i], HDR_BG)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        bg = row_bg(ri, row) if row_bg else None
        for ci, val in enumerate(row):
            _cell_text(cells[ci], val, mono=(ci in mono_cols))
            if bg: _shade(cells[ci], bg)
    if widths:
        for ri in range(len(t.rows)):
            for ci, w in enumerate(widths):
                t.rows[ri].cells[ci].width = Inches(w)
    doc.add_paragraph()
    return t


def callout(doc, title, body, color=AMBER, bg=NEW_BG):
    """A single-cell shaded table — the only box style Docs imports reliably."""
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.rows[0].cells[0]; _shade(c, bg)
    c.text = ""
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = color
    for para_text in ([body] if isinstance(body, str) else body):
        q = c.add_paragraph(); q.paragraph_format.space_after = Pt(3)
        rr = q.add_run(para_text); rr.font.size = Pt(10)
    doc.add_paragraph()
    return t


def template_box(doc, label, lines, note=None):
    """A blank template slot for someone to fill in later."""
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.rows[0].cells[0]; _shade(c, BOX_BG); c.text = ""
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label); r.bold = True; r.font.size = Pt(10); r.font.color.rgb = BLUE
    if note:
        q = c.add_paragraph(); q.paragraph_format.space_after = Pt(6)
        rr = q.add_run(note); rr.font.size = Pt(9); rr.italic = True; rr.font.color.rgb = MUTED
    for ln in lines:
        q = c.add_paragraph(); q.paragraph_format.space_after = Pt(3)
        rr = q.add_run(ln); rr.font.size = Pt(10); rr.font.name = "Consolas"
        rr.font.color.rgb = MUTED
    doc.add_paragraph()
    return t


def bullets(doc, items, style="List Bullet"):
    for it in items:
        if isinstance(it, tuple):
            p = doc.add_paragraph(style=style)
            r = p.add_run(it[0]); r.bold = True
            p.add_run(" — " + it[1])
        else:
            doc.add_paragraph(it, style=style)


def steps(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Number")


def trim_png(png_bytes, pad=12):
    """Crop the white margin off a headless-Chrome screenshot.

    The renderer shoots a fixed 2400px-tall viewport, so a short diagram arrives sitting on
    ~1500px of blank white. Uncropped, every image lands 15in tall in Word — nearly two
    pages of mostly nothing. Crop to the ink, then pad slightly.
    """
    import io
    from PIL import Image, ImageChops
    im = Image.open(io.BytesIO(png_bytes)).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    box = ImageChops.difference(im, bg).getbbox()
    if box:
        l, t, r, b = box
        l = max(0, l - pad); t = max(0, t - pad)
        r = min(im.width, r + pad); b = min(im.height, b + pad)
        im = im.crop((l, t, r, b))
    out = io.BytesIO(); im.save(out, format="PNG")
    return out.getvalue(), im.width, im.height


def image(doc, png_bytes, width_in=6.6, caption=None, max_h_in=8.4):
    """Insert a diagram, cropped to its ink and capped so it never exceeds one page."""
    import io
    png_bytes, w, h = trim_png(png_bytes)
    # honour width_in, but if that would run past the page, scale by height instead
    if h / w * width_in > max_h_in:
        width_in = max_h_in * w / h
    doc.add_picture(io.BytesIO(png_bytes), width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption); r.font.size = Pt(9); r.italic = True; r.font.color.rgb = MUTED


def rule(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "CCCCCC")
    pbdr.append(bottom); pPr.append(pbdr)
